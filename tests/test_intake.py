"""The intake rules accept clean files, reject each planted problem, and reproduce
fjordvik-data exactly when the December files are loaded."""

import json
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO))
from intake_rules import COLUMNS, REPORTING_MONTH, Reference, accepted_rows, missing_submissions, validate  # noqa: E402

EXPECTED = json.loads((REPO / "expected" / "expected_results.json").read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def ref():
    return Reference()


def read_submission(path):
    """Read the workbook the way the flow does: only the rows inside the table."""
    wb = load_workbook(path)
    ws = wb["Submission"]
    table = ws.tables["Submission"]
    cells = list(ws[table.ref])
    columns = [c.value for c in cells[0]]
    rows = [dict(zip(columns, [c.value for c in r])) for r in cells[1:]]
    return columns, rows


def site_of(path):
    return path.name.split("_")[0]


@pytest.mark.parametrize("path", sorted((REPO / "intake" / "clean").glob("*.xlsx")), ids=lambda p: p.name)
def test_clean_files_are_accepted(path, ref):
    columns, rows = read_submission(path)
    assert validate(rows, columns, site_of(path), ref) == []


def test_problem_files_fail_exactly_the_planned_rules(ref):
    planned = {"FI01": ["R5", "R9"], "FI03": ["R8"], "SE01": ["R7"], "FI02": [], "EE01": []}
    for path in sorted((REPO / "intake" / "problems").glob("*.xlsx")):
        columns, rows = read_submission(path)
        got = sorted({f["rule"] for f in validate(rows, columns, site_of(path), ref)},
                     key=lambda r: int(r[1:]))
        assert got == planned[site_of(path)], path.name
    assert not (REPO / "intake" / "problems" / f"EE02_{REPORTING_MONTH}.xlsx").exists()


def test_loading_the_clean_files_reproduces_fjordvik_data(ref):
    """Central list before the run + accepted December rows == the full shared dataset."""
    central = [{k: r[k] for k in ("date_key", "site_id", "source_id")}
               | {"activity_amount": float(r["activity_amount"]), "unit": r["unit"],
                  "emission_factor": float(r["emission_factor"]), "tco2e": float(r["tco2e"])}
               for r in ref.central_before()]
    for r in central:
        r["date_key"] = int(r["date_key"])
    for path in sorted((REPO / "intake" / "clean").glob("*.xlsx")):
        _cols, rows = read_submission(path)
        central += accepted_rows(rows, site_of(path), ref)
    key = lambda r: (r["date_key"], r["site_id"], r["source_id"])
    source = sorted(({"date_key": int(r["date_key"]), "site_id": r["site_id"], "source_id": r["source_id"],
                      "activity_amount": float(r["activity_amount"]), "unit": r["unit"],
                      "emission_factor": float(r["emission_factor"]), "tco2e": float(r["tco2e"])}
                     for r in ref.facts), key=key)
    assert sorted(central, key=key) == source
    assert len(source) == 1422


def test_a_second_load_of_the_same_month_is_refused():
    """R10 makes the flow safe to rerun: a file already loaded is never added twice."""
    ref = Reference(loaded_rows=Reference().facts)  # December already loaded
    columns, rows = read_submission(REPO / "intake" / "clean" / f"FI02_{REPORTING_MONTH}.xlsx")
    assert [f["rule"] for f in validate(rows, columns, "FI02", ref)] == ["R10"]


@pytest.mark.parametrize("change, rule", [
    (lambda rows: [{**r, "month": "2025-11"} for r in rows], "R3"),
    (lambda rows: [{**r, "site_id": "FI01"} for r in rows], "R2"),
    (lambda rows: rows + [{**rows[0], "source_id": "S01"}], "R4"),
    (lambda rows: [{**rows[0], "activity_amount": "n/a"}] + rows[1:], "R6"),
    (lambda rows: [{**rows[0], "activity_amount": -5.0}] + rows[1:], "R6"),
    # Triple electricity (S05); refrigerant (S04, the first row) is exempt from R9.
    (lambda rows: [{**r, "activity_amount": r["activity_amount"] * 3} if r["source_id"] == "S05" else r
                   for r in rows], "R9"),
])
def test_each_remaining_rule_catches_its_problem(change, rule, ref):
    """FI02 (head office) reports no natural gas, so S01 there trips R4."""
    columns, rows = read_submission(REPO / "intake" / "clean" / f"FI02_{REPORTING_MONTH}.xlsx")
    assert rule in {f["rule"] for f in validate(change(rows), columns, "FI02", ref)}


def test_wrong_columns_stop_everything_else(ref):
    columns, rows = read_submission(REPO / "intake" / "clean" / f"FI02_{REPORTING_MONTH}.xlsx")
    bad = ["site", *COLUMNS[1:]]
    assert [f["rule"] for f in validate(rows, bad, "FI02", ref)] == ["R1"]


def test_refrigerant_is_exempt_from_the_change_check(ref):
    """Top-ups are zero most months; a ratio against last year would reject good data."""
    columns, rows = read_submission(REPO / "intake" / "clean" / f"FI01_{REPORTING_MONTH}.xlsx")
    rows = [{**r, "activity_amount": 40.0} if r["source_id"] == "S04" else r for r in rows]
    assert validate(rows, columns, "FI01", ref) == []


def test_missing_submissions_lists_sites_without_an_accepted_file():
    assert EXPECTED["missing_submissions"]["clean"] == []
    # Rejected files must be fixed and sent again, so they count as not yet in.
    assert EXPECTED["missing_submissions"]["problems"] == ["EE02", "FI01", "FI03", "SE01"]


def test_expected_results_are_current(ref):
    """expected/expected_results.json is what the reference rules give today."""
    for kind in ("clean", "problems"):
        for name, exp in EXPECTED[kind].items():
            path = REPO / "intake" / kind / name
            columns, rows = read_submission(path)
            fails = validate(rows, columns, site_of(path), ref)
            assert fails == exp["failures"], name


def test_every_workbook_holds_one_named_table():
    for path in list((REPO / "intake").rglob("*.xlsx")) + list((REPO / "sharepoint").rglob("*.xlsx")):
        wb = load_workbook(path)
        tables = [t for ws in wb.worksheets for t in ws.tables.values()]
        assert len(tables) == 1, path.name


def test_seed_emissions_stop_at_november_2025():
    wb = load_workbook(REPO / "sharepoint" / "seed" / "Emissions.xlsx")
    keys = [row[0] for row in wb["Emissions"].iter_rows(min_row=2, values_only=True)]
    assert max(keys) == 20251101 and len(keys) == 1422 - 43

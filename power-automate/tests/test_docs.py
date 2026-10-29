"""The documents agree with the rules and the seed files."""

import sys
from pathlib import Path

from openpyxl import load_workbook

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO))
from intake_rules import CHANGE_HIGH, CHANGE_LOW, RULES  # noqa: E402

DOCS = REPO / "docs"


def test_every_rule_is_documented():
    text = (DOCS / "intake-rules.md").read_text(encoding="utf-8")
    for rule in RULES:
        assert f"| {rule} |" in text, rule
    assert CHANGE_LOW == 0.5 and CHANGE_HIGH == 2.0, "update 'half to double' in the docs"


def test_documented_seed_row_counts_match_the_workbooks():
    text = (DOCS / "sharepoint-lists.md").read_text(encoding="utf-8")
    expected = {"Sites": 6, "Sources": 8, "EmissionFactors": 24, "Emissions": 1379}
    for name, rows in expected.items():
        wb = load_workbook(REPO / "sharepoint" / "seed" / f"{name}.xlsx")
        assert wb.active.max_row - 1 == rows, name
        assert f"{rows:,}" in text, name


def test_seed_columns_match_the_documented_lists():
    text = (DOCS / "sharepoint-lists.md").read_text(encoding="utf-8")
    for name in ("Sites", "Sources", "EmissionFactors", "Emissions"):
        wb = load_workbook(REPO / "sharepoint" / "seed" / f"{name}.xlsx")
        header = [c.value for c in next(wb.active.iter_rows(max_row=1))]
        line = next(l for l in text.splitlines() if l.startswith(f"| {name} |"))
        for column in header:
            assert f"{column} (" in line, f"{name}.{column}"
